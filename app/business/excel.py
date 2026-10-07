"""SCH-09 엑셀 내보내기 · SCH-10 엑셀 가져오기 (미리보기 → 매핑 수정 → 반영)"""
import io
import re
import uuid
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from sqlalchemy.orm import Session

from app.business import audit, members, schedules
from app.business.common import ApiException, not_found, unprocessable
from app.domain.model import Duty
from app.persistence.db import now
from app.persistence.models import ImportPreview

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLORS = {Duty.D: "FFF2CC", Duty.E: "E2EFDA", Duty.N: "DDEBF7", Duty.O: "EDEDED", Duty.AL: "FCE4D6", Duty.ED: "E4DFEC"}
# 엑셀 코드 → 듀티 제안값. 나머지는 사용자가 매핑
SUMMARY_ROW = re.compile(r"[DEN] 인원 \(필요 \d+\)")  # 내보내기 파일의 인원 합계 행
ALIASES = {"OFF": Duty.O, "/": Duty.O, "휴": Duty.O, "DAY": Duty.D, "EVE": Duty.E, "NIGHT": Duty.N}


def export_schedule(db: Session, user_id: uuid.UUID, schedule_id: uuid.UUID) -> tuple[bytes, str]:
    m = members.require_head(db, user_id)
    s = schedules.find(db, m, schedule_id)
    r = schedules.roster(db, s)
    days = r.days()
    wb = Workbook()
    ws = wb.active
    ws.title = s.year_month
    ws.append(["이름"] + [d.day for d in days] + ["D", "E", "N", "OFF", "AL"])
    for n in r.nurses:
        row = [r.cell(n.id, d) for d in days]
        ws.append([n.name] + [str(x) if x else None for x in row] + [row.count(x) for x in (Duty.D, Duty.E, Duty.N,
                                                                                            Duty.O, Duty.AL)])
        for i, x in enumerate(row):
            if x:
                c = ws.cell(ws.max_row, 2 + i)
                c.fill = PatternFill("solid", fgColor=COLORS[x])
                c.alignment = Alignment(horizontal="center")
    for duty in (Duty.D, Duty.E, Duty.N):
        ws.append([f"{duty} 인원 (필요 {r.rules.required(duty)})"]
                  + [sum(1 for n in r.nurses if r.cell(n.id, d) == duty and n.affiliated(d)) for d in days])
    ws.column_dimensions["A"].width = 18
    # 내보내기는 정보 유출 경로이므로 감사 로그 필수
    audit.log(db, m.ward_id, user_id, "SCHEDULE_EXPORTED", "SCHEDULE", s.id)
    return _save(wb), f"schedule-{s.year_month}.xlsx"


def create_preview(db: Session, user_id: uuid.UUID, schedule_id: uuid.UUID, data: bytes) -> dict:
    """
    행=간호사(첫 열 이름), 열=날짜(첫 행 일자). 근무표에는 아직 반영하지 않는다.
    동명이인·미등록 이름은 자동 매칭하지 않고 nurseMappings로 지정받는다. 보호 셀(AL·소속 기간 밖)은 가져오지 않는다
    """
    m = members.require_head(db, user_id)
    s = schedules.find(db, m, schedule_id)
    r = schedules.roster(db, s)
    by_name: dict[str, list] = {}
    for n in r.nurses:
        by_name.setdefault(n.name.strip(), []).append(n.id)
    ws = _open(data)
    header = next(ws.iter_rows(min_row=ws.min_row, max_row=ws.min_row), None)
    day_cols = {}
    for c in header or []:
        t = _text(c.value).removesuffix("일")
        if t.isdigit() and 1 <= int(t) <= len(r.days()):
            day_cols[c.column] = r.month.replace(day=int(t))
    if not day_cols:
        raise unprocessable("INVALID_EXCEL_FORMAT", "첫 행에서 날짜 열을 찾지 못했습니다")
    rows, cells, nurse_map, duty_map = [], [], {}, {}
    for row in ws.iter_rows(min_row=ws.min_row + 1):
        name = _text(row[0].value) if row else ""
        if not name or SUMMARY_ROW.fullmatch(name):
            continue
        row_no = row[0].row
        candidates = by_name.get(name, [])
        rows.append({"rowNumber": row_no, "name": name, "candidates": [str(x) for x in candidates]})
        if len(candidates) == 1:
            nurse_map[str(row_no)] = str(candidates[0])
        for c in row:
            # AL은 연차 승인으로만 생기므로 가져오지 않는다 (기존 AL 셀은 반영 시에도 보호됨)
            if c.column in day_cols and (raw := _text(c.value).upper()) and raw != Duty.AL:
                cells.append({"rowNumber": row_no, "date": day_cols[c.column].isoformat(), "raw": raw})
                if raw not in duty_map and (suggested := _suggest(raw)) is not None:
                    duty_map[raw] = suggested
    p = ImportPreview(ward_id=m.ward_id, schedule_id=s.id, rows=rows, cells=cells, nurse_mappings=nurse_map,
                      duty_mappings=duty_map, created_by=user_id)
    db.add(p)
    db.flush()
    return _view(p)


def update_preview(db: Session, user_id: uuid.UUID, schedule_id: uuid.UUID, preview_id: uuid.UUID,
                   nurse_mappings: list[tuple[int, uuid.UUID | None]] | None,
                   duty_mappings: list[tuple[str, Duty | None]] | None) -> dict:
    """nurseId=null은 그 행을 건너뜀, dutyCode=null은 미배정으로 반영"""
    m = members.require_head(db, user_id)
    p = _preview(db, m, schedule_id, preview_id)
    rows = {r["rowNumber"] for r in p.rows}
    nurse_ids = {n.id for n in schedules.roster(db, schedules.find(db, m, schedule_id)).nurses}
    nm, dm = dict(p.nurse_mappings), dict(p.duty_mappings)
    for row_no, nurse_id in nurse_mappings or []:
        if row_no not in rows or (nurse_id is not None and nurse_id not in nurse_ids):
            raise unprocessable("INVALID_MAPPING", f"매핑할 수 없는 행 또는 간호사입니다: {row_no}")
        nm[str(row_no)] = str(nurse_id) if nurse_id else None
    codes = {c["raw"] for c in p.cells}
    for raw, duty in duty_mappings or []:
        if raw not in codes or duty == Duty.AL:
            raise unprocessable("INVALID_MAPPING", f"매핑할 수 없는 코드입니다: {raw} (AL은 연차 승인으로만)")
        dm[raw] = duty
    p.nurse_mappings, p.duty_mappings = nm, dm
    return _view(p)


def apply_preview(db: Session, user_id: uuid.UUID, schedule_id: uuid.UUID, preview_id: uuid.UUID, base_version: int,
                  lock_token: str | None) -> dict:
    m = members.require_head(db, user_id)
    p = _preview(db, m, schedule_id, preview_id)
    unresolved = _unresolved(p)
    if unresolved["rows"] or unresolved["codes"]:
        raise unprocessable("UNRESOLVED_MAPPING", "매핑되지 않은 행 또는 코드가 있습니다", unresolved)
    s = schedules.find(db, m, schedule_id)
    schedules.editable(db, s, m, lock_token, base_version)
    r = schedules.roster(db, s)
    nurses = {n.id: n for n in r.nurses}
    changes, seen = [], set()
    for c in p.cells:
        nurse_id = p.nurse_mappings.get(str(c["rowNumber"]))
        if nurse_id is None:
            continue
        nurse_id, d = uuid.UUID(nurse_id), date.fromisoformat(c["date"])
        # 보호 셀(AL·소속 기간 밖)은 그대로 둔다. 같은 간호사가 여러 행에 있으면 첫 행만
        if (nurse_id, d) in seen or r.cell(nurse_id, d) == Duty.AL or not nurses[nurse_id].affiliated(d):
            continue
        seen.add((nurse_id, d))
        duty = p.duty_mappings[c["raw"]]
        changes.append((nurse_id, d, Duty(duty) if duty else None))
    changed = schedules.apply_changes(db, s, changes)
    p.applied_at = now()
    audit.log(db, m.ward_id, user_id, "SCHEDULE_IMPORTED", "SCHEDULE", s.id, None, f"{changed} cells")
    return {"schedule": schedules.view(db, s, m), "violations": schedules.check(db, s)}


def _preview(db: Session, m: members.Member, schedule_id: uuid.UUID, preview_id: uuid.UUID) -> ImportPreview:
    p = db.get(ImportPreview, preview_id)
    if p is None or p.ward_id != m.ward_id or p.schedule_id != schedule_id:
        raise not_found("PREVIEW_NOT_FOUND")
    if p.applied_at is not None:
        raise ApiException(409, "INVALID_RESOURCE_STATE", "이미 반영된 미리보기입니다")
    return p


def _unresolved(p: ImportPreview) -> dict:
    return {"rows": [r["rowNumber"] for r in p.rows if str(r["rowNumber"]) not in p.nurse_mappings],
            "codes": sorted({c["raw"] for c in p.cells} - set(p.duty_mappings))}


def _view(p: ImportPreview) -> dict:
    return {"id": p.id, "scheduleId": p.schedule_id, "status": "APPLIED" if p.applied_at else "READY",
            "rows": p.rows, "cells": p.cells,
            "nurseMappings": [{"rowNumber": int(k), "nurseId": v} for k, v in sorted(p.nurse_mappings.items(),
                                                                                    key=lambda x: int(x[0]))],
            "dutyMappings": [{"raw": k, "dutyCode": v} for k, v in sorted(p.duty_mappings.items())],
            "unresolved": _unresolved(p), "createdAt": p.created_at.isoformat() + "Z"}


def _suggest(raw: str) -> Duty | None:
    if raw in ALIASES:
        return ALIASES[raw]
    try:
        d = Duty(raw)
    except ValueError:
        return None
    return None if d == Duty.AL else d  # AL은 연차 승인으로만 (COM-03)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _open(data: bytes):
    try:
        return load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    except Exception:  # 깨진 파일·xlsx가 아닌 파일은 종류가 다양한 예외로 실패한다
        raise unprocessable("INVALID_EXCEL_FORMAT", "엑셀 형식을 읽을 수 없습니다") from None


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
