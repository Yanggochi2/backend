"""API-SCH-09~12 엑셀 내보내기·가져오기"""
import io

from openpyxl import load_workbook

from conftest import NEXT, YM, add_nurses, err, fill_rotation, member, ward_head


def _xlsx(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_then_import_with_mapping():
    h = ward_head()
    add_nurses(h, 4)
    s = fill_rotation(h, h.data("POST", "/wards/me/schedules", {"yearMonth": YM}))
    url = f"/wards/me/schedules/{s['id']}"
    x = h("GET", url + "/export.xlsx")
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.active
    assert ws.cell(1, 1).value == "이름" and ws.cell(2, 1).value == s["nurses"][0]["name"]
    ws.cell(2, 3).value = "OFF"  # 첫 간호사 2일 → O
    ws.cell(3, 4).value = "???"  # 모르는 코드 → 매핑 필요
    ws.cell(4, 1).value = "없는사람"  # 미등록 이름 → 행 매핑 필요

    err(h("POST", url + "/import-previews", files={"file": ("x.xlsx", b"nope")}), 422, "INVALID_EXCEL_FORMAT")
    p = h.data("POST", url + "/import-previews", files={"file": ("s.xlsx", _xlsx(wb))}, status=202)
    assert p["unresolved"] == {"rows": [4], "codes": ["???"]}
    purl = f"{url}/import-previews/{p['id']}"
    err(h("POST", purl + "/apply", {"baseVersion": s["version"]}), 422, "UNRESOLVED_MAPPING")
    err(h("PATCH", purl, {"dutyMappings": [{"raw": "???", "dutyCode": "AL"}]}), 422, "INVALID_MAPPING")
    p = h.data("PATCH", purl, {"dutyMappings": [{"raw": "???", "dutyCode": "O"}],
                               "nurseMappings": [{"rowNumber": 4, "nurseId": None}]})
    assert p["unresolved"] == {"rows": [], "codes": []}
    err(h("POST", purl + "/apply", {"baseVersion": s["version"] - 1}), 409, "VERSION_CONFLICT")
    r = h.data("POST", purl + "/apply", {"baseVersion": s["version"]})
    first = s["nurses"][0]["id"]
    cell = next(c for c in r["schedule"]["cells"] if c["nurseId"] == first and c["date"] == NEXT.replace(day=2)
                .isoformat())
    assert cell["dutyCode"] == "O" and r["schedule"]["version"] == s["version"] + 1
    err(h("POST", purl + "/apply", {"baseVersion": r["schedule"]["version"]}), 409, "INVALID_RESOURCE_STATE")


def test_export_is_head_only():
    h = ward_head()
    n, _ = member(h)
    s = h.data("POST", "/wards/me/schedules", {"yearMonth": YM})
    err(n("GET", f"/wards/me/schedules/{s['id']}/export.xlsx"), 403, "FORBIDDEN")
